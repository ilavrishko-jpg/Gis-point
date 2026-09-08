#!/usr/bin/env python3
"""Звіряє вшиту в index.html базу з вихідними вивантаженнями.

Запуск:  python3 test/check_data.py <тека з .xlsx>
Виходить з кодом 1, якщо хоч одне поле не збіглося.
"""
import openpyxl, json, sys, datetime, pathlib, collections

NAME = "Ім'я"
ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else ".")

def cell(v):
    if v is None: return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%d.%m.%Y") if (v.hour or v.minute) == 0 else v.strftime("%d.%m.%Y %H:%M")
    if isinstance(v, datetime.date): return v.strftime("%d.%m.%Y")
    return str(v).strip()

def read(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [cell(h) for h in rows[0]]
    return [{hdr[i]: cell(r[i]) for i in range(len(hdr))}
            for r in rows[1:] if any(c not in (None, "") for c in r)]

html = (ROOT / "index.html").read_text(encoding="utf-8")
base = json.loads(html.split('<script type="application/json" id="base-data">')[1]
                      .split("</script>")[0].replace("\\u003c", "<"))
rows = base["rows"]

it = read(SRC / "3e2ef633-MSPO_IT_unit.xlsx")
fb = read(SRC / "82a92221-MSPO_Flyby.xlsx")
emb_it = [r for r in rows if r["__source"] == "ІТ-аутсорсинг"]
emb_fb = [r for r in rows if r["__source"] == "FlyBy"]

bad = []
if len(rows) != len(it) + len(fb):
    bad.append("кількість: %d вшито, %d у файлах" % (len(rows), len(it) + len(fb)))

IT_PAIRS = [(NAME, NAME), ("Посада", "Посада"), ("Стенд", "Стенд"), ("Країна", "Країна"),
            ("Email", "Email"), ("LinkedIn URL", "LinkedIn URL"), ("Розмір компанії", "Розмір компанії"),
            ("Збагачення", "Збагачення"), ("Повідомлення (останнє)", "Повідомлення (останнє)"),
            ("Ресурс", "Ресурс"), ("Статус", "Статус")]
FB_PAIRS = [("Компанія", "Компанія"), ("Стенд", "Стенд"), ("Контакт", NAME), ("Посада", "Посада"),
            ("Канал", "Канал"), ("Збагачення", "Збагачення"), ("Статус", "Статус"),
            ("Останнє повідомлення клієнта", "Повідомлення (останнє)")]

for src, emb, pairs, tag in ((it, emb_it, IT_PAIRS, "IT"), (fb, emb_fb, FB_PAIRS, "FB")):
    for s, e in zip(src, emb):
        who = s.get(NAME) or s.get("Контакт") or s.get("Компанія") or "?"
        for sc, ec in pairs:
            if s.get(sc, "") and s.get(sc, "") != e.get(ec, ""):
                bad.append("%s [%s] %s: файл=%r апка=%r" % (tag, who, sc, s.get(sc), e.get(ec)))
        if s.get("Компанія") and s["Компанія"] != e["Компанія"]:
            bad.append("%s [%s] Компанія: %r -> %r" % (tag, who, s["Компанія"], e["Компанія"]))

lost = [s for s, e in zip(it, emb_it)
        if s.get("Коментар") and not (e["Гачок"] or e["Сегмент"] or e["Пріоритет"])]
if lost: bad.append("розібрано не всі коментарі: %d втрачено" % len(lost))

meet_src = sum(1 for s in it if s.get("Дата зустрічі")) + sum(1 for s in fb if s.get("Час зустрічі"))
meet_emb = sum(1 for r in rows if r["Зустріч"])
if meet_src != meet_emb: bad.append("зустрічі: %d у файлах, %d в апці" % (meet_src, meet_emb))

print("вшито: %d рядків, %d колонок" % (len(rows), len(base["columns"])))
print("за джерелом:", dict(collections.Counter(r["__source"] for r in rows)))
print("зустрічей:", meet_emb)
print("назву компанії відновлено:", dict(collections.Counter(
    r["Компанія визначена"] for r in rows if r["Компанія визначена"])))
print("без компанії й без пошти:", sum(1 for r in rows if "(без компанії)" in r["Компанія"]))

if bad:
    print("\nРОЗБІЖНОСТІ (%d):" % len(bad))
    for x in bad[:25]: print("  !", x)
    sys.exit(1)
print("\nусі поля збіглися з вихідними файлами")
