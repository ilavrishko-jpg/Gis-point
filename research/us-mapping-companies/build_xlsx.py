#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Assemble all research CSVs into one multi-tab .xlsx workbook."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(OUT, "US_Mapping_Companies.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")

def style_table(ws, nrows, ncols, header_row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if nrows >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{nrows}"

def add_csv_sheet(wb, title, csv_name, widths=None):
    ws = wb.create_sheet(title)
    path = os.path.join(OUT, csv_name)
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            if r > 1:
                cell.alignment = WRAP
    if rows:
        style_table(ws, len(rows), len(rows[0]))
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    return ws

wb = Workbook()

# ---- 1. Overview ----
ov = wb.active
ov.title = "Огляд"
ov.column_dimensions["A"].width = 42
ov.column_dimensions["B"].width = 80
ov["A1"] = "Компанії США: lidar / aerial / mobile mapping, tomography, cartography"
ov["A1"].font = TITLE_FONT
overview_rows = [
    ("", ""),
    ("Підготовлено для", "GIS-Point"),
    ("Дата", "червень 2026"),
    ("Категорії", "lidar mapping · aerial mapping · mobile mapping · tomography (GPR/SUE/геофізика) · cartography/GIS"),
    ("Цільові розмірні групи", "① 30–100 · ② 200–300 · ③ 300–500 працівників"),
    ("", ""),
    ("ЩО ВСЕРЕДИНІ (вкладки)", ""),
    ("Компанії (138)", "Курований реєстр ідентифікованих фірм США: штат, тип, працівники, розмір, категорії, достовірність"),
    ("Матриця штат×розмір", "К-сть HQ-фірм по штату × розмірній групі"),
    ("Великі фірми", "17 великих AEC/geospatial: точні розміри (SEC 10-K) + geospatial-підрозділ + офіси + штати"),
    ("Присутність по штатах", "Скільки великих фірм має офіс у кожному штаті (усі 50 + DC)"),
    ("Підсумок розміри / категорії", "Агреговані лічильники"),
    ("Галузь і джерела", "Знаменник Census/IBISWorld + джерела + застереження"),
    ("", ""),
    ("КЛЮЧОВІ ЦИФРИ", ""),
    ("Усього фірм у реєстрі", "138"),
    ("у групі 30–100", "37"),
    ("у групі 200–300", "9"),
    ("у групі 300–500", "8"),
    ("Галузь NAICS 541370 (усього бізнесів)", "≈17 511 (IBISWorld 2025); ≈7 382 establishments-роботодавців (Census 2020)"),
    ("Середній розмір фірми в галузі", "3,8 працівника — переважна більшість мікрофірми"),
    ("Покриття штатів офісами великих фірм", "усі 50 штатів + DC"),
    ("", ""),
    ("ВАЖЛИВЕ ЗАСТЕРЕЖЕННЯ", "Це вибірка ідентифікованих фірм, а не повний перепис. Кількість працівників — переважно оцінки (±30–50%); див. колонку Confidence. Точний розподіл Census по штатах×розмірних класах недоступний (домени census.gov/bls.gov заблоковані політикою середовища)."),
]
for i, (a, b) in enumerate(overview_rows, start=2):
    ov.cell(row=i, column=1, value=a).font = Font(bold=bool(a) and a.isupper())
    cell = ov.cell(row=i, column=2, value=b)
    cell.alignment = WRAP

# ---- data sheets ----
add_csv_sheet(wb, "Компанії (138)", "companies.csv",
              widths=[46, 10, 14, 18, 12, 26, 16, 12])
add_csv_sheet(wb, "Матриця штат×розмір", "state_size_matrix.csv",
              widths=[10] + [9]*6 + [9, 16])
add_csv_sheet(wb, "Великі фірми", "large_firms.csv",
              widths=[40, 9, 16, 14, 18, 30, 11, 60, 11])
add_csv_sheet(wb, "Присутність по штатах", "large_firm_state_presence.csv",
              widths=[8, 14, 100])
add_csv_sheet(wb, "Підсумок розміри", "summary_by_bucket.csv",
              widths=[16, 16, 16])
add_csv_sheet(wb, "Підсумок категорії", "summary_by_category.csv",
              widths=[18] + [9]*6 + [9])

# ---- Industry & sources ----
src = wb.create_sheet("Галузь і джерела")
src.column_dimensions["A"].width = 50
src.column_dimensions["B"].width = 90
src["A1"] = "Галузевий знаменник, джерела та застереження"
src["A1"].font = TITLE_FONT
src_rows = [
    ("", ""),
    ("ГАЛУЗЬ NAICS 541370", ""),
    ("Бізнесів (вкл. самозайнятих), IBISWorld 2025", "17 511"),
    ("Establishments-роботодавців, Census CBP 2020", "≈7 382"),
    ("Зайнятість у галузі (роботодавці), 2023", "≈55 206"),
    ("Обсяг ринку (дохід), 2024–25", "$11,3–11,5 млрд"),
    ("Середній розмір фірми", "3,8 працівника"),
    ("Зайнятість по штатах (BLS QCEW 2023)", "TX 9 139 · FL 6 684 · CA 4 085"),
    ("Концентрація (location quotient)", "AK 2,38 · FL 1,92 · TX 1,85 · LA 1,81"),
    ("", ""),
    ("ДЖЕРЕЛА", ""),
    ("IBISWorld — Surveying & Mapping Services (NAICS 541370)", "ibisworld.com/united-states/.../surveying-mapping-services/1407/"),
    ("US Census Bureau — County Business Patterns / SUSB", "census.gov/programs-surveys/cbp.html (заблоковано в середовищі)"),
    ("BLS QCEW — зайнятість у галузі 2023", "bls.gov/opub/ted/2024/... (заблоковано)"),
    ("MAPPS — асоціація приватних geospatial-фірм (~160)", "mapps.org"),
    ("Профілі компаній", "LinkedIn, ZoomInfo, RocketReach, PitchBook, Crunchbase, D&B, сайти фірм, SEC 10-K (NVEE, BWMN)"),
    ("", ""),
    ("ЗАСТЕРЕЖЕННЯ", ""),
    ("Достовірність HQ-штатів", "висока"),
    ("Достовірність к-сті працівників", "переважно med/low — діапазони агрегаторів, ±30–50%"),
    ("AEC-фірми", "показано розмір усієї фірми; geospatial-підрозділ зазвичай менший (50–800 осіб)"),
    ("Точний Census по штатах×розмірних класах", "недоступний: домени census.gov/api.census.gov/bls.gov заблоковані egress-політикою; републікатори за платним доступом"),
    ("Не включено (не-US)", "Eagle Mapping (CA), North West Geomatics (CA), GeoFly (DE), SUMO (UK)"),
    ("Не включено (історичні/поглинуті)", "AeroMetric, Photo Science, HJW, Quantum Spatial (=NV5), 3001 International, Terrapoint"),
]
for i, (a, b) in enumerate(src_rows, start=2):
    src.cell(row=i, column=1, value=a).font = Font(bold=bool(a) and a.isupper())
    src.cell(row=i, column=2, value=b).alignment = WRAP

wb.save(XLSX)
print("Wrote", XLSX)
print("Sheets:", wb.sheetnames)
print("Size bytes:", os.path.getsize(XLSX))
