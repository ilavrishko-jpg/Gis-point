# -*- coding: utf-8 -*-
import re
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
wb=load_workbook("/home/user/Gis-point/market-analysis/unit-economics-model.xlsx")
C={}
REF=re.compile(r"(?:'([^']+)'!)?\$?([A-Z]{1,2})\$?(\d+)")
RNG=re.compile(r"(?:'([^']+)'!)?\$?([A-Z]{1,2})\$?(\d+):\$?([A-Z]{1,2})\$?(\d+)")
def val(sh,col,row,d=0):
    k=(sh,col,row)
    if k in C: return C[k]
    if d>60: raise RecursionError(k)
    c=wb[sh].cell(row,ci(col)).value
    v=ev(c[1:],sh,d+1) if isinstance(c,str) and c.startswith("=") else (float(c) if isinstance(c,(int,float)) else 0.0)
    C[k]=v; return v
def ev(f,cur,d=0):
    f=RNG.sub(lambda m:"["+",".join(repr(val(m.group(1) or cur,gl(c),r,d)) for c in range(ci(m.group(2)),ci(m.group(4))+1) for r in range(int(m.group(3)),int(m.group(5))+1))+"]",f)
    f=REF.sub(lambda m:repr(val(m.group(1) or cur,m.group(2),int(m.group(3)),d)),f)
    f=re.sub(r"\bSUM\(","__s(",f); f=re.sub(r"\bSUMPRODUCT\(","__sp(",f); f=re.sub(r"\bIF\(","__if(",f)
    f=re.sub(r"(?<![<>=!])=(?!=)","==",f).replace("^","**")
    return float(eval(f,{"__builtins__":{}},{"__s":lambda *a:sum(x for g in a for x in (g if isinstance(g,list) else [g])),
        "__sp":lambda a,b:sum(x*y for x,y in zip(a,b)),"__if":lambda c,a,b:a if c else b}))
def find(sh,frag,col=1):
    for r in range(1,140):
        v=wb[sh].cell(r,col).value
        if isinstance(v,str) and frag.lower() in v.lower(): return r
    raise KeyError(f"{sh}:{frag}")
OK=[0,0]
def chk(l,sh,frag,col,exp,tol=.02,occ=0):
    rows=[r for r in range(1,140) if isinstance(wb[sh].cell(r,1).value,str) and frag.lower() in wb[sh].cell(r,1).value.lower()]
    r=rows[occ]; v=val(sh,col,r)
    ok=abs(v-exp)<=max(abs(exp)*tol,.51); OK[0]+=ok; OK[1]+=1
    print(f"  {'OK  ' if ok else 'FAIL'} {l:<52}книга={v:>12,.1f}  очік.={exp:>12,.1f}".replace(","," "))
U="2.Юніт-економіка"; P="5.P&L 2026-2029"; S="6.Сценарії"; M="3.Ринки"; D="4.Матриця рішення"
print("A. Геодезія @ 15 000 грн/день — «важка» (B) проти «легкої» (C)")
chk("Фіксовані витрати, важка",U,"Фіксовані витрати","B",2_103_000); chk("Фіксовані витрати, легка",U,"Фіксовані витрати","C",1_512_000)
chk("Беззбитковість, днів — важка",U,"Точка беззбитковості, днів","B",155); chk("Беззбитковість, днів — легка",U,"Точка беззбитковості, днів","C",110)
chk("Беззбитковість % util — важка",U,"% utilization","B",0.672); chk("Беззбитковість % util — легка",U,"% utilization","C",0.477)
chk("@62% ВП — важка (ЗБИТОК)",U,"@62%: ВАЛОВИЙ","B",-163_640); chk("@62% ВП — легка",U,"@62%: ВАЛОВИЙ","C",455_880)
chk("@75% ВП — важка",U,"@75%: ВАЛОВИЙ","B",243_000); chk("@75% ВП — легка",U,"@75%: ВАЛОВИЙ","C",868_500)
chk("Різниця структур при 62%",U,"Різниця «легка»","B",619_520)
print("B. Геологія @ 70 000 грн/день — днів до беззбитковості (стеля 210)")
for rate,e1,e2,e3 in [(40000,99,77,64),(55000,69,53,45),(70000,53,41,34),(85000,43,33,28)]:
    r=[x for x in range(1,140) if wb[U].cell(x,1).value==rate][0]
    for col,exp,nm in (("C",e1,"повна"),("D",e2,"самортизована"),("E",e3,"мінімальна")):
        v=val(U,col,r); ok=abs(v-exp)<=max(exp*.03,.51); OK[0]+=ok; OK[1]+=1
        print(f"  {'OK  ' if ok else 'FAIL'} {f'@{rate} грн/день, {nm}':<52}книга={v:>12,.0f}  очік.={exp:>12,.0f}".replace(","," "))
print("B2. Геологія — результат за завантаженням")
for u_,rev,gp,gm,pf in [(0.45,6_615_000,2_357_700,.356,14_554),(0.55,8_085_000,3_546_300,.439,21_891),
                        (0.70,10_290_000,5_329_200,.518,32_896),(0.80,11_760_000,6_517_800,.554,40_233)]:
    r=[x for x in range(1,140) if isinstance(wb[U].cell(x,1).value,float) and abs(wb[U].cell(x,1).value-u_)<1e-9][0]
    for col,exp,nm in (("C",rev,"виручка"),("E",gp,"ВАЛОВИЙ"),("G",pf,"ВП/FTE $")):
        v=val(U,col,r); ok=abs(v-exp)<=max(abs(exp)*.02,.51); OK[0]+=ok; OK[1]+=1
        print(f"  {'OK  ' if ok else 'FAIL'} {f'util {u_:.0%}: {nm}':<52}книга={v:>12,.0f}  очік.={exp:>12,.0f}".replace(","," "))
print("C. IT")
chk("Повна вартість інженера, $/рік",U,"Повна вартість інженера","B",51_528)
chk("Собівартість години, $",U,"Собівартість 1 оплачуваної","B",35.3)
chk("Внутрішній T&M $26: ВП/інженера",U,"Внутрішній UA, T&M","E",-13_594)
print("D. Підписки")
chk("D1 LTV",U,"D1 SaaS","B",91_000); chk("D1 LTV/CAC",U,"D1 SaaS","C",10.7); chk("D1 payback",U,"D1 SaaS","D",9.3)
chk("D2 LTV/CAC",U,"D2 DaaS","C",13.3); chk("D2 payback",U,"D2 DaaS","D",4.5)
chk("D3 LTV/CAC (провальний)",U,"D3 Агро","C",1.8); chk("D3 payback",U,"D3 Агро","D",30.5)
print("Ринки")
chk("M2 геологія ВП 2029",M,"M2","J",42.6,.03); chk("M1 геодезія ВП 2029",M,"M1","J",7.5,.05)
chk("M6 мережі ВП 2029",M,"M6","J",49.6,.03)
chk("РАЗОМ SOM 2029",M,"РАЗОМ","H",335,.02); chk("РАЗОМ FTE",M,"РАЗОМ","K",117,.02)
print("Матриця")
chk("M6 бал",D,"M6","G",4.60); chk("M2 геологія бал (тепер ядро)",D,"M2","G",4.45); chk("M1 геодезія бал",D,"M1","G",2.60)
print("P&L")
chk("База 2026: виручка",P,"РАЗОМ","C",63.5,.02,0); chk("База 2026: валовий прибуток",P,"РАЗОМ","E",18.0,.02,0)
chk("Ціль 2029: виручка",P,"РАЗОМ","C",229.9,.02,1); chk("Ціль 2029: валовий прибуток",P,"РАЗОМ","E",124.9,.02,1)
chk("Ціль 2029: ВП на 1 FTE",P,"РАЗОМ","F",24_243,.02,1)
chk("База 2026: EBITDA (ПРИБУТОК)",P,"EBITDA","C",7.8,.05,0); chk("Ціль 2029: EBITDA",P,"EBITDA","C",62.8,.03,1)
print("Сценарії")
for f,e in [("A. Війна",230),("B. Перемир",425),("C. Ескалац",115),("ОЧІКУВАНЕ",256)]:
    chk(f"{f}: виручка 2029",S,f,"D",e,.02)
print(f"\n{'='*76}\nРЕЗУЛЬТАТ: {OK[0]}/{OK[1]} перевірок пройдено")
